# -*- coding: utf-8 -*-
import openpyxl
import copy

PATH = r"C:\Users\Sergio Gomez\Downloads\UAO\SISTEMAS OPERATIVOSD\ComandosUnixAGOSTO25_2026.xlsx"

# row -> (col_index_1based, new_value)   col 4=Solaris
FIXES = [
    (27, 4, "SI (verificado en vivo: Solaris 11.4 incluye gawk)"),
    (42, 4, "SI (verificado en vivo: locate esta disponible)"),
    (61, 4, "SI (verificado en vivo: ps aux funciona directo, sin necesitar /usr/ucb/ps)"),
    (62, 4, "SI (verificado en vivo: ps ax funciona directo)"),
    (63, 4, "NO (verificado en vivo: pstree no esta instalado por defecto)"),
    (69, 4, "NO (verificado en vivo: Solaris usa 'swap -s' / 'swap -l', no 'swapinfo'; swapinfo es de FreeBSD/HP-UX)"),
]

wb = openpyxl.load_workbook(PATH)
ws = wb.active
template_style = copy.copy(ws.cell(row=3, column=4)._style)

for row, col, val in FIXES:
    cell = ws.cell(row=row, column=col)
    cell.value = val
    cell._style = copy.copy(template_style)

wb.save(PATH)
print("Fixed", len(FIXES), "Solaris cells with live-verified data")
